"""Tests for pure file analysis-capability classification."""

from __future__ import annotations

import io

from core.analysis_capability import (
    AnalysisCapability,
    assess_capability,
    is_supported_type,
    not_checked,
)
from core.deliverable_discovery_extraction import extract_units


def _pdf_bytes(pages: list[str]) -> bytes:
    import fitz

    document = fitz.open()
    for text in pages:
        page = document.new_page()
        if text:
            page.insert_text((72, 72), text)
    data = document.tobytes()
    document.close()
    return data


def _docx_bytes(text: str) -> bytes:
    from docx import Document

    document = Document()
    document.add_paragraph(text)
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def _xlsx_bytes() -> bytes:
    from openpyxl import Workbook

    workbook = Workbook()
    worksheet = workbook.active
    worksheet.append(["Item", "Quantity", "Price"])
    worksheet.append(["Pump", 2, 1500])
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def test_readable_pdf() -> None:
    data = _pdf_bytes(
        ["Submit this complete technical proposal with all required schedules before closing time."]
    )

    result = assess_capability("instructions.pdf", data)

    assert result.capability is AnalysisCapability.READABLE
    assert result.label == "Can be analysed"
    assert result.type_label == "PDF"
    assert result.readable_parts == 1
    assert result.unreadable_parts == 0


def test_partly_readable_pdf_reports_empty_page() -> None:
    data = _pdf_bytes(
        [
            "Submit this complete technical proposal with all required schedules "
            "before closing time.",
            "",
        ]
    )

    result = assess_capability("instructions.pdf", data)

    assert result.capability is AnalysisCapability.READABLE
    assert result.label == "Can be analysed — 1 of 2 page(s) have no readable text"
    assert result.readable_parts == 1
    assert result.unreadable_parts == 1


def test_empty_pdf_has_no_readable_text() -> None:
    result = assess_capability("empty.pdf", _pdf_bytes([""]))

    assert result.capability is AnalysisCapability.NO_READABLE_TEXT
    assert result.label == "No readable text found — OCR may be needed"
    assert result.unreadable_parts == 1


def test_corrupt_pdf_could_not_be_opened() -> None:
    result = assess_capability("x.pdf", b"not a pdf")

    assert result.capability is AnalysisCapability.COULD_NOT_OPEN
    assert result.label == "Could not be opened — the file may be damaged or password-protected"


def test_unsupported_file_does_not_require_valid_data() -> None:
    result = assess_capability("model.dwg", b"")

    assert result.capability is AnalysisCapability.UNSUPPORTED
    assert result.type_label == "DWG file"
    assert result.label == (
        "Stored only — this file type cannot be analysed (PDF, Word and Excel only)"
    )


def test_word_and_excel_files_are_readable() -> None:
    word = assess_capability("scope.docx", _docx_bytes("The contractor must submit a schedule."))
    excel = assess_capability("prices.xlsx", _xlsx_bytes())

    assert word.capability is AnalysisCapability.READABLE
    assert word.type_label == "Word document"
    assert excel.capability is AnalysisCapability.READABLE
    assert excel.type_label == "Excel workbook"


def test_not_checked_distinguishes_supported_and_unsupported_files() -> None:
    assert not_checked("a.pdf").capability is AnalysisCapability.NOT_CHECKED
    assert not_checked("a.dwg").capability is AnalysisCapability.UNSUPPORTED


def test_supported_type_is_case_insensitive() -> None:
    for filename in ("a.pdf", "a.PDF", "a.docx", "a.DOCX", "a.xlsx", "a.XLSX"):
        assert is_supported_type(filename)
    for filename in ("a.dwg", "a.txt"):
        assert not is_supported_type(filename)


def test_empty_pdf_extraction_reason_is_honest() -> None:
    units = extract_units("empty.pdf", None, io.BytesIO(_pdf_bytes([""])))

    assert units[0].failure_reason is not None
    assert "OCR may be needed" in units[0].failure_reason
    assert "scanned" not in units[0].failure_reason.lower()
