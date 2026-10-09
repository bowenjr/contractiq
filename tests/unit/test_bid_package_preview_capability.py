"""Capability checks for mutation-free Bid package previews."""

from __future__ import annotations

import asyncio
import io
import os
from collections.abc import Iterator
from pathlib import Path
from types import ModuleType

import pytest

from core.analysis_capability import AnalysisCapability
from core.bid_package_storage import UnsafeIntakeSourceError
from scripts.asgi_acceptance_ops08 import request
from tests.unit.test_bid_package_storage import BID_ID, _service
from tests.unit.test_ops11_ui import BID_ID as UI_BID_ID
from tests.unit.test_ops11_ui import ui_app as _ops11_ui_app


@pytest.fixture
def capability_ui_app(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> Iterator[ModuleType]:
    """Reuse the configured-inbox application fixture from the OPS-11 route tests."""
    yield from _ops11_ui_app.__wrapped__(tmp_path, monkeypatch)


def _pdf_bytes(texts: list[str]) -> bytes:
    import fitz

    document = fitz.open()
    for text in texts:
        page = document.new_page()
        if text:
            page.insert_text((72, 72), text)
    data = document.tobytes()
    document.close()
    return data


def _docx_bytes(text: str) -> bytes:
    from docx import Document

    document = Document()
    document.add_paragraph(text)
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def _xlsx_bytes() -> bytes:
    from openpyxl import Workbook

    workbook = Workbook()
    worksheet = workbook.active
    worksheet.append(["Item", "Quantity", "Price"])
    worksheet.append(["Pump", 2, 1500])
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def _file_bytes(root: Path) -> dict[str, bytes]:
    return {
        path.relative_to(root).as_posix(): path.read_bytes()
        for path in root.rglob("*")
        if path.is_file()
    }


def _tree_state(root: Path) -> dict[str, bytes | None]:
    return {
        path.relative_to(root).as_posix(): path.read_bytes() if path.is_file() else None
        for path in root.rglob("*")
    }


def _table_counts(database_path: str) -> dict[str, int]:
    import sqlite3

    with sqlite3.connect(database_path) as connection:
        tables = [
            str(row[0])
            for row in connection.execute(
                "SELECT name FROM sqlite_master WHERE type='table' ORDER BY name"
            )
        ]
        return {
            table: int(connection.execute(f'SELECT count(*) FROM "{table}"').fetchone()[0])
            for table in tables
        }


def test_read_preview_file_is_bounded_safe_and_mutation_free(tmp_path: Path) -> None:
    service, _repository, _database, source, _managed = _service(tmp_path)
    package = source / "package"
    package.mkdir()
    normal = package / "normal.pdf"
    normal.write_bytes(b"normal bytes")
    (package / "large.pdf").write_bytes(b"larger than limit")
    outside = tmp_path / "outside.pdf"
    outside.write_bytes(b"outside")
    os.symlink(outside, package / "linked.pdf")
    before = _file_bytes(source)

    assert (
        service.storage.read_preview_file(
            "landing", "package", "normal.pdf", expected_bid_id=BID_ID, max_bytes=100
        )
        == b"normal bytes"
    )
    assert (
        service.storage.read_preview_file(
            "landing", "package", "large.pdf", expected_bid_id=BID_ID, max_bytes=4
        )
        is None
    )
    for unsafe_path in ("../x", str(normal.resolve()), "linked.pdf"):
        with pytest.raises(UnsafeIntakeSourceError):
            service.storage.read_preview_file(
                "landing",
                "package",
                unsafe_path,
                expected_bid_id=BID_ID,
                max_bytes=100,
            )

    assert _file_bytes(source) == before


def test_preview_capabilities_classifies_files_without_mutation(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    service, _repository, database, source, managed = _service(tmp_path)
    readable = _pdf_bytes(
        [
            "Submit the complete signed technical proposal and every required schedule "
            "before closing."
        ]
    )
    empty = _pdf_bytes([""])
    word = _docx_bytes("The contractor must submit a complete construction schedule.")
    excel = _xlsx_bytes()
    files = {
        "a-readable.pdf": readable,
        "b-empty.pdf": empty,
        "c-model.dwg": b"drawing bytes",
        "d-corrupt.pdf": b"not a pdf",
        "e-scope.docx": word,
        "f-prices.xlsx": excel,
    }
    max_checked_bytes = max(len(data) for data in files.values())
    files["g-oversized.pdf"] = b"x" * (max_checked_bytes + 1)
    for name, data in files.items():
        (source / name).write_bytes(data)
    preview = service.preview(BID_ID, "landing")
    counts_before = _table_counts(database.db_path)
    managed_before = _tree_state(managed)
    source_before = _file_bytes(source)
    original_read = service.storage.read_preview_file
    read_paths: list[str] = []

    def guarded_read(
        source_root_alias: str,
        source_folder: str | None,
        relative_path: str,
        *,
        expected_bid_id: str | None = None,
        max_bytes: int,
    ) -> bytes | None:
        if relative_path.endswith(".dwg"):
            raise AssertionError("unsupported files must not be read")
        read_paths.append(relative_path)
        return original_read(
            source_root_alias,
            source_folder,
            relative_path,
            expected_bid_id=expected_bid_id,
            max_bytes=max_bytes,
        )

    monkeypatch.setattr(service.storage, "read_preview_file", guarded_read)

    capabilities = service.preview_capabilities(
        BID_ID,
        preview,
        max_file_bytes=max_checked_bytes,
    )

    assert capabilities["a-readable.pdf"].capability is AnalysisCapability.READABLE
    assert capabilities["b-empty.pdf"].capability is AnalysisCapability.NO_READABLE_TEXT
    assert capabilities["b-empty.pdf"].label == "No readable text found — OCR may be needed"
    assert capabilities["c-model.dwg"].capability is AnalysisCapability.UNSUPPORTED
    assert capabilities["d-corrupt.pdf"].capability is AnalysisCapability.COULD_NOT_OPEN
    assert capabilities["e-scope.docx"].capability is AnalysisCapability.READABLE
    assert capabilities["f-prices.xlsx"].capability is AnalysisCapability.READABLE
    assert capabilities["g-oversized.pdf"].capability is AnalysisCapability.NOT_CHECKED
    assert "c-model.dwg" not in read_paths
    assert _table_counts(database.db_path) == counts_before
    assert _tree_state(managed) == managed_before
    assert _file_bytes(source) == source_before
    assert service.preview(BID_ID, "landing").source_fingerprint == preview.source_fingerprint


def test_preview_capabilities_stops_supported_reads_at_time_budget(tmp_path: Path) -> None:
    service, _repository, _database, source, _managed = _service(tmp_path)
    (source / "a-first.pdf").write_bytes(
        _pdf_bytes(
            ["Submit the complete signed technical proposal and every required schedule today."]
        )
    )
    (source / "b-model.dwg").write_bytes(b"drawing")
    (source / "c-later.docx").write_bytes(_docx_bytes("This file should not be opened."))
    preview = service.preview(BID_ID, "landing")
    clock_values = iter((0.0, 0.0, 9.0))

    def clock() -> float:
        return next(clock_values)

    capabilities = service.preview_capabilities(
        BID_ID,
        preview,
        time_budget_seconds=8.0,
        clock=clock,
    )

    assert capabilities["a-first.pdf"].capability is AnalysisCapability.READABLE
    assert capabilities["b-model.dwg"].capability is AnalysisCapability.UNSUPPORTED
    assert capabilities["c-later.docx"].capability is AnalysisCapability.NOT_CHECKED


def test_preview_route_displays_capability_without_writes_or_model_calls(
    capability_ui_app: ModuleType, monkeypatch: pytest.MonkeyPatch
) -> None:
    ui_app = capability_ui_app

    async def scenario() -> None:
        package = ui_app.INTAKE_SOURCE_ROOTS["inbox_1"] / "Initial-Package"
        (package / "readable.pdf").write_bytes(
            _pdf_bytes(
                ["Submit the complete signed technical proposal and every required schedule today."]
            )
        )
        (package / "empty.pdf").write_bytes(_pdf_bytes([""]))
        (package / "model.dwg").write_bytes(b"synthetic drawing")
        audit_before = len(ui_app.bid_repository.list_audit(UI_BID_ID))

        def model_call_forbidden(*_args: object, **_kwargs: object) -> object:
            raise AssertionError("preview must not call a model endpoint")

        monkeypatch.setattr(ui_app.llm_client, "chat", model_call_forbidden)
        monkeypatch.setattr(ui_app.discovery_llm_client, "chat", model_call_forbidden)
        status, _, body = await request(
            ui_app.app,
            "POST",
            f"/bids/{UI_BID_ID}/package-intake-addenda/preview",
            {
                "release_type": "INITIAL_PACKAGE",
                "source_root_alias": "inbox_1",
                "source_folder": "Initial-Package",
                "intake_mode": "initial",
                "exact_customer_reference": "RFP-CAPABILITY",
                "received_channel": "EMAIL",
            },
        )
        page = body.decode()

        assert status == 200
        assert "Can it be analysed?" in page
        assert "No readable text found — OCR may be needed" in page
        assert "Stored only — this file type cannot be analysed (PDF, Word and Excel only)" in page
        assert "Analysis:" in page
        assert "scanned image" not in page.lower()
        assert len(ui_app.bid_repository.list_audit(UI_BID_ID)) == audit_before

    asyncio.run(scenario())


def test_intake_preview_display_without_capabilities_keeps_previous_keys(
    capability_ui_app: ModuleType,
) -> None:
    ui_app = capability_ui_app
    preview = ui_app.bid_package_service.preview(UI_BID_ID, "inbox_1", "Initial-Package")

    display = ui_app._intake_preview_display(preview, "synthetic-token")

    assert {
        "token",
        "folder_name",
        "file_count",
        "total_bytes",
        "directories",
        "files",
        "warnings",
    } <= display.keys()
    assert display["token"] == "synthetic-token"
    assert len(display["files"]) == preview.file_count
    assert all(
        {
            "original_relative_path",
            "original_filename",
            "extension",
            "detected_media_type",
            "byte_size",
            "sha256",
            "is_archive",
            "is_duplicate",
            "needs_identification",
        }
        <= row.keys()
        for row in display["files"]
    )
